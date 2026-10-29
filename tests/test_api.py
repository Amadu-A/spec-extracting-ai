# tests/test_api.py
"""Полный API-путь без браузера: PDF→worker→SSE replay→строгий JSON→равнозначный Excel."""

import asyncio
import json
from datetime import timedelta
from io import BytesIO

import pymupdf
import pytest
from openpyxl import load_workbook

from app.domain.errors import ServiceError
from app.domain.models import JobState, utcnow
from tests.conftest import KEY_A, KEY_B, PASSWORD

HEADERS = {"X-API-Key": KEY_A}


async def create(client, data, pages="1-2", key=None):
    """Создаёт настоящее асинхронное задание через multipart публичного API."""
    headers = {**HEADERS, **({"Idempotency-Key": key} if key else {})}
    return await client.post(
        "/api/v1/extractions",
        headers=headers,
        files={"file": ("sample.pdf", data, "application/pdf")},
        data={"pages": pages, "project_code": "TEST-PROJECT"},
    )


async def test_headless_e2e_json_excel_sse(root, client, pdf_bytes):
    """Весь pipeline работает без UI; экспорт и SSE берутся из committed state."""
    response = await create(client, pdf_bytes)
    assert response.status_code == 202
    status = response.json()
    job_id = status["job_id"]
    assert status["status"] == "queued" and not root.model.calls
    before = await client.get(f"/api/v1/extractions/{job_id}/result.json", headers=HEADERS)
    assert before.status_code == 409
    await root.processor.execute(job_id)
    status = (await client.get(f"/api/v1/extractions/{job_id}", headers=HEADERS)).json()
    assert status["status"] == "completed" and status["progress"] == 100
    assert status["processed_pages"] == 2 and status["item_count"] == 2
    result = await client.get(status["links"]["json"], headers=HEADERS)
    assert result.status_code == 200
    payload = result.json()
    assert list(payload) == ["items"]
    assert [x["Pages"] for x in payload["items"]] == ["1", "2"]
    assert all(x["Quantity"] == "2,5" for x in payload["items"])
    excel = await client.get(status["links"]["xlsx"], headers=HEADERS)
    book = load_workbook(BytesIO(excel.content))
    rows = list(book["Спецификация"].values)
    assert list(rows[0]) == list(payload["items"][0])
    assert [list(row) for row in rows[1:]] == [list(row.values()) for row in payload["items"]]
    events = await client.get(status["links"]["events"], headers=HEADERS)
    assert "event: page_started" in events.text and "event: text_delta" in events.text
    assert "event: item_validated" in events.text and "event: job_completed" in events.text
    ids = [int(line[4:]) for line in events.text.splitlines() if line.startswith("id: ")]
    assert ids == list(range(1, status["last_event_id"] + 1))
    resumed = await client.get(
        status["links"]["events"], headers={**HEADERS, "Last-Event-ID": str(ids[-2])}
    )
    assert resumed.text.count("id: ") == 1
    snapshot = (await client.get(f"/api/v1/extractions/{job_id}/snapshot", headers=HEADERS)).json()
    assert snapshot["partial"] == payload
    assert snapshot["last_event_id"] == ids[-1]
    assert all(call["image_count"] >= 1 for call in root.model.calls)
    stages = [
        stage.name for page in (await root.repository.get(job_id)).pages for stage in page.stages
    ]
    assert {"Извлечение текста", "Рендер страницы", "Экспорт JSON", "Экспорт Excel"}.issubset(
        stages
    )


async def test_idempotency_conflict_and_isolation(root, client, pdf_bytes):
    """Одинаковый ключ возвращает тот же job; чужой credential не раскрывает его данные."""
    first, repeated = (
        await create(client, pdf_bytes, key="stable"),
        await create(client, pdf_bytes, key="stable"),
    )
    assert first.json()["job_id"] == repeated.json()["job_id"]
    conflict = await create(client, pdf_bytes, pages="3", key="stable")
    assert (
        conflict.status_code == 409 and conflict.json()["error"]["code"] == "idempotency_conflict"
    )
    identifier = first.json()["job_id"]
    for suffix in [
        "",
        "/events",
        "/result.json",
        "/result.xlsx",
        "/warnings",
        "/pages/1/preview",
        "/snapshot",
    ]:
        other = await client.get(
            f"/api/v1/extractions/{identifier}{suffix}", headers={"X-API-Key": KEY_B}
        )
        assert other.status_code == 404
    assert (await client.get(f"/api/v1/extractions/{identifier}")).status_code == 401


async def test_retry_invalidates_and_avoids_duplicates(root, client, pdf_bytes):
    """Повтор листа создаёт ревизию и заменяет зависимый суффикс без удвоения строк."""
    identifier = (await create(client, pdf_bytes, pages="1-3")).json()["job_id"]
    await root.processor.execute(identifier)
    retried = await client.post(f"/api/v1/extractions/{identifier}/pages/2/retry", headers=HEADERS)
    assert retried.status_code == 202 and retried.json()["revision"] == 1
    assert retried.json()["item_count"] == 0
    assert (
        await client.get(f"/api/v1/extractions/{identifier}/result.json", headers=HEADERS)
    ).status_code == 409
    await root.processor.execute(identifier)
    result = await client.get(f"/api/v1/extractions/{identifier}/result.json", headers=HEADERS)
    assert len(result.json()["items"]) == 3
    invalid = await client.post(f"/api/v1/extractions/{identifier}/pages/4/retry", headers=HEADERS)
    assert invalid.status_code == 422


async def test_cancellation_of_queued_and_running(root, client, pdf_bytes):
    """Отмена идемпотентна и завершает ожидающий model stream без успеха экспорта."""
    identifier = (await create(client, pdf_bytes)).json()["job_id"]
    for _ in range(2):
        cancelled = await client.post(f"/api/v1/extractions/{identifier}/cancel", headers=HEADERS)
        assert cancelled.json()["status"] == "cancelled"
    await root.processor.execute(identifier)
    assert not root.model.calls
    identifier = (await create(client, pdf_bytes)).json()["job_id"]
    root.model.delay = 1
    task = asyncio.create_task(root.processor.execute(identifier))
    for _ in range(200):
        if root.model.calls:
            break
        await asyncio.sleep(0.02)
    await client.post(f"/api/v1/extractions/{identifier}/cancel", headers=HEADERS)
    await asyncio.wait_for(task, 4)
    status = (await client.get(f"/api/v1/extractions/{identifier}", headers=HEADERS)).json()
    assert status["status"] == "cancelled" and status["progress"] < 100
    assert (
        await client.get(f"/api/v1/extractions/{identifier}/result.json", headers=HEADERS)
    ).status_code == 409


@pytest.mark.parametrize(
    ("failure", "code"),
    [
        ("{bad", "invalid_model_json"),
        (ServiceError("vlm_unavailable", "Модель недоступна", 503), "vlm_unavailable"),
    ],
)
async def test_model_failure_is_not_empty_success(root, client, pdf_bytes, failure, code):
    """Недоступная модель/плохой JSON дают failed, а не успешный пустой items."""
    root.model.responses = [failure]
    identifier = (await create(client, pdf_bytes)).json()["job_id"]
    await root.processor.execute(identifier)
    status = (await client.get(f"/api/v1/extractions/{identifier}", headers=HEADERS)).json()
    assert status["status"] == "failed" and status["error_code"] == code
    warnings = (
        await client.get(f"/api/v1/extractions/{identifier}/warnings", headers=HEADERS)
    ).json()
    assert warnings["warnings"][0]["code"] == code


async def test_prompt_snapshot_api_isolation_csrf_ttl(root, client, pdf_bytes):
    """UI prompt влияет только на новые задания текущей сессии; API всегда получает default."""
    session = (await client.get("/api/v1/session")).json()
    login = await client.post(
        "/api/v1/session/login",
        headers={"X-CSRF-Token": session["csrf"]},
        json={"password": PASSWORD},
    )
    assert login.status_code == 200
    assert (
        "httponly" in login.headers["set-cookie"].lower()
        and "max-age" not in login.headers["set-cookie"].lower()
    )
    csrf = {"X-CSRF-Token": login.json()["csrf"]}
    denied = await client.put("/api/v1/session/prompt", json={"prompt": "Custom"})
    assert denied.status_code == 403
    denied = await client.put(
        "/api/v1/session/prompt",
        headers={**csrf, "Origin": "https://evil.example"},
        json={"prompt": "Custom"},
    )
    assert denied.status_code == 403
    assert (
        await client.put("/api/v1/session/prompt", headers=csrf, json={"prompt": "Custom"})
    ).status_code == 200
    ui = await client.post(
        "/api/v1/extractions",
        headers=csrf,
        files={"file": ("a.pdf", pdf_bytes, "application/pdf")},
        data={"pages": "1", "project_code": "P"},
    )
    ui_job = await root.repository.get(ui.json()["job_id"])
    assert ui_job.prompt == "Custom"
    api_job = await root.repository.get((await create(client, pdf_bytes)).json()["job_id"])
    assert api_job.prompt == root.extractions.default_prompt
    restored = await client.post("/api/v1/session/prompt/restore", headers=csrf)
    assert (
        restored.json()["is_default"]
        and restored.json()["prompt"] == root.extractions.default_prompt
    )
    assert (await root.repository.get(ui_job.id)).prompt == "Custom"
    session = await root.sessions.resolve(client.cookies["spec_session"])
    session.expires_at = utcnow() - timedelta(seconds=1)
    await root.repository.save_session(session)
    assert (await client.get("/api/v1/session/prompt")).status_code == 401


async def test_pdf_security_and_preview(root, client, pdf_bytes):
    """Проверяются сигнатура/MIME/пароль/повреждение, UUID paths и ограничение preview."""
    for data, code in [(b"not pdf", "unsupported_file"), (b"%PDF-1.7\ninvalid", "corrupted_pdf")]:
        result = await create(client, data)
        assert result.status_code in (415, 422) and result.json()["error"]["code"] == code
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        protected = document.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret"
        )
    result = await create(client, protected)
    assert result.json()["error"]["code"] == "password_pdf"
    identifier = (await create(client, pdf_bytes)).json()["job_id"]
    image = await client.get(
        f"/api/v1/extractions/{identifier}/pages/3/preview?width=300", headers=HEADERS
    )
    assert image.status_code == 200 and image.content.startswith(b"\x89PNG")
    oversized = await client.get(
        f"/api/v1/extractions/{identifier}/pages/1/preview?width=90000", headers=HEADERS
    )
    assert oversized.status_code == 422
    with pytest.raises(ServiceError):
        root.storage.path("../../outside")


async def test_ready_and_export_warning(root, client, pdf_bytes):
    """Readiness отличается от live; непроверенная строка видима и исключена из экспорта."""
    assert (await client.get("/api/v1/health/ready")).status_code == 200
    root.model.available = False
    assert (await client.get("/api/v1/health/ready")).status_code == 503
    assert (await client.get("/api/v1/health/live")).status_code == 200
    root.model.responses = [
        json.dumps(
            {
                "rows": [
                    {
                        "row_id": "1",
                        "designation": "",
                        "manufacturer": "",
                        "caption": "Неясная позиция",
                        "units": "",
                        "quantity": "",
                        "verified": False,
                        "continuation": False,
                        "continues_next": False,
                    }
                ],
                "group_context": [],
                "warnings": ["Плохая читаемость"],
                "completeness": "uncertain",
            }
        )
    ]
    identifier = (await create(client, pdf_bytes, pages="1")).json()["job_id"]
    await root.processor.execute(identifier)
    job = await root.repository.get(identifier)
    assert job.state == JobState.WARNINGS and not job.export().items
    result = await client.get(f"/api/v1/extractions/{identifier}/result.xlsx", headers=HEADERS)
    assert "Предупреждения" in load_workbook(BytesIO(result.content)).sheetnames


async def test_body_limit_before_multipart_parser(root, pdf_bytes):
    """Content-Length и chunked body ограничиваются до создания большого upload spool."""
    import httpx

    from app.presentation.http.app import create_app

    root.settings.max_upload_mb = 1
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app(root)), base_url="http://test"
    ) as connection:
        response = await connection.post(
            "/api/v1/documents",
            headers={**HEADERS, "Content-Length": str(3 * 1024 * 1024)},
            content=b"x",
        )
        assert response.status_code == 413
        assert response.json()["error"]["code"] == "upload_limit"

        async def chunks():
            """Отправляет multipart без Content-Length, превышая лимит фактического потока."""
            yield b'--limit-test\r\nContent-Disposition: form-data; name="file"; filename="x.pdf"\r\nContent-Type: application/pdf\r\n\r\n'
            for _ in range(3):
                yield b"x" * (1024 * 1024)
            yield b"\r\n--limit-test--\r\n"

        response = await connection.post(
            "/api/v1/documents",
            headers={**HEADERS, "Content-Type": "multipart/form-data; boundary=limit-test"},
            content=chunks(),
        )
        assert response.status_code == 413 and response.json()["error"]["code"] == "upload_limit"
