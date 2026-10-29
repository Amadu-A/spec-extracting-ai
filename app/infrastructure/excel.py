# app/infrastructure/excel.py
"""Excel строится из того же Export DTO, что JSON, без повторного вызова модели."""

import asyncio
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.domain.models import Export, Warning

COLUMNS = ("Caption", "Units", "Quantity", "UserComment1", "Pages")


class ExcelExporter:
    """Выполняет CPU/ZIP операции openpyxl вне HTTP event loop."""

    async def excel(self, result: Export, warnings: list[Warning]) -> bytes:
        """Реализует application-порт без зависимости от FastAPI."""
        return await asyncio.to_thread(build_excel, result, warnings)


def build_excel(result: Export, warnings: list[Warning]) -> bytes:
    """Сохраняет все ячейки строками, включая формулы и десятичные количества.

    data_type=s предотвращает formula injection без изменения экспортного значения.
    Вспомогательный лист предупреждений явно отделён от спецификации.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "Спецификация"
    sheet.append(COLUMNS)
    for item in result.items:
        sheet.append([getattr(item, name) for name in COLUMNS])
        for cell in sheet[sheet.max_row]:
            cell.data_type, cell.number_format = "s", "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087BFA")
    for column, width in zip("ABCDE", [100, 15, 18, 30, 12], strict=True):
        sheet.column_dimensions[column].width = width
    if warnings:
        diagnostic = book.create_sheet("Предупреждения")
        diagnostic.append(["PDF-страница", "Код", "Причина", "Непроверенная позиция"])
        for warning in warnings:
            diagnostic.append(
                [
                    str(warning.pdf_page),
                    warning.code,
                    warning.message,
                    warning.row.caption if warning.row else "",
                ]
            )
            for cell in diagnostic[diagnostic.max_row]:
                cell.data_type = "s"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
